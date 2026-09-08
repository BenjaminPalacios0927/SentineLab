import sys
import os
from datetime import datetime, timedelta, date
from typing import Dict, Optional
from PyQt6.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout,
                             QHBoxLayout, QComboBox, QPushButton, QLabel, QDialog,
                             QScrollArea, QFrame, QFileDialog, QMessageBox, QDateEdit,
                             QGridLayout, QSizePolicy)
from PyQt6.QtCore import QTimer, Qt, QDate, pyqtSignal, QThread
from PyQt6.QtGui import QPainter, QColor, QFont


def ocultar_consola_windows():
    if os.name == 'nt':
        try:
            import ctypes
            ctypes.windll.kernel32.FreeConsole()
        except Exception:
            pass

try:
    from interface_graphics import InterfaceGraphics
    GRAPHICS_UTILITY_AVAILABLE = True
except ImportError:
    GRAPHICS_UTILITY_AVAILABLE = False

try:
    from attendance_processor import ProcesadorReporteAsistencia
    PROCESSOR_AVAILABLE = True
except ImportError:
    PROCESSOR_AVAILABLE = False


# =========================================================
# ESTILOS (idénticos al original)
# =========================================================
STYLE_BTN_PRIMARY = """
    QPushButton {
        background-color: rgb(15, 110, 86);
        color: white;
        font-family: Arial;
        font-size: 17px;
        font-weight: bold;
        padding: 14px 32px;
        border-radius: 10px;
        border: none;
    }
    QPushButton:hover { background-color: rgb(10, 80, 62); }
    QPushButton:disabled { background-color: rgb(159, 225, 203); color: white; }
"""
STYLE_BTN_GHOST = """
    QPushButton {
        background-color: white;
        color: rgb(24, 95, 165);
        font-family: Arial;
        font-size: 17px;
        font-weight: bold;
        padding: 14px 32px;
        border-radius: 10px;
        border: 2px solid rgb(24, 95, 165);
    }
    QPushButton:hover { background-color: rgb(230, 241, 251); }
    QPushButton:disabled { background-color: rgb(241, 239, 232); color: rgb(181, 212, 244); }
"""
STYLE_DATE = """
    QDateEdit {
        background-color: rgb(241, 239, 232);
        color: rgb(26, 26, 26);
        font-family: Arial;
        font-size: 16px;
        padding: 10px 12px;
        border-radius: 8px;
        border: 1px solid rgb(181, 212, 244);
        min-width: 160px;
    }
    QDateEdit:focus { background-color: white; border: 2px solid rgb(24, 95, 165); }
    QCalendarWidget QWidget { background-color: white; color: black; }
    QCalendarWidget QToolButton { color: black; background-color: transparent; icon-size: 20px; }
    QCalendarWidget QMenu { background-color: white; color: black; }
    QCalendarWidget QSpinBox { background-color: white; color: black; selection-background-color: #2196F3; }
    QCalendarWidget QAbstractItemView:enabled {
        color: black; background-color: white;
        selection-background-color: #2196F3; selection-color: white;
    }
"""
STYLE_COMBO = """
    QComboBox {
        background-color: rgb(241, 239, 232);
        color: rgb(26, 26, 26);
        font-family: Arial;
        font-size: 17px;
        padding: 10px 14px;
        border-radius: 8px;
        border: 1px solid rgb(181, 212, 244);
        min-width: 420px;
    }
    QComboBox:focus { background-color: white; border: 2px solid rgb(24, 95, 165); }
    QComboBox::drop-down { border: none; }
    QComboBox QAbstractItemView {
        background-color: white; color: rgb(26, 26, 26);
        selection-background-color: rgb(24, 95, 165); selection-color: white;
        font-size: 17px;
    }
"""


# =========================================================
# HILOS
# =========================================================
class HiloCargaEmpleados(QThread):
    """
    Carga la lista de empleados desde device_users (sin filtro de fechas).
    Esto garantiza que el combo siempre muestre todos los empleados
    registrados en el lector, independientemente de si tienen asistencia
    en el rango de los últimos 60 días.
    """
    completado = pyqtSignal(object)
    error      = pyqtSignal(str)

    def __init__(self, procesador):
        super().__init__()
        self.procesador = procesador

    def run(self):
        try:
            # Sin rango de fechas → pre-pobla empleados desde device_users
            self.procesador.cargar_datos_desde_railway()
            self.completado.emit(self.procesador.obtener_lista_empleados())
        except Exception as e:
            self.error.emit(str(e))


class HiloGenerarInforme(QThread):
    """
    Recarga los datos del período exacto desde Railway y filtra por empleado.
    Hace la recarga dentro del hilo para no bloquear el hilo principal.
    """
    completado = pyqtSignal(object)
    error      = pyqtSignal(str)

    def __init__(self, procesador, id_empleado, fecha_inicio, fecha_fin):
        super().__init__()
        self.procesador  = procesador
        self.id_empleado = id_empleado
        self.f_inicio    = fecha_inicio
        self.f_fin       = fecha_fin

    def run(self):
        try:
            # Recargar desde Railway con el rango exacto del informe
            # para que filtrar_por_periodo tenga los registros del período
            self.procesador.cargar_datos_desde_railway(
                self.f_inicio.date(), self.f_fin.date())
            resultado = self.procesador.filtrar_por_periodo(
                self.id_empleado, self.f_inicio, self.f_fin)
            self.completado.emit(resultado)
        except Exception as e:
            self.error.emit(str(e))


# =========================================================
# CANVAS DE FONDO
# =========================================================
class CanvasFondo(QWidget):
    def __init__(self, parent, graphics_engine):
        super().__init__(parent)
        self.tiempo   = 0
        self.graphics = graphics_engine
        self.setAttribute(Qt.WidgetAttribute.WA_TranslucentBackground, False)

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        if self.graphics:
            self.graphics.dibujar_gradiente_animado(
                painter, self.width(), self.height(), self.tiempo)
        else:
            painter.fillRect(0, 0, self.width(), self.height(), QColor(30, 60, 114))
        self.tiempo += 1


# =========================================================
# FUNCIÓN PURA: EVALUACIÓN
# =========================================================
def _generar_evaluacion(datos: dict) -> str:
    faltas      = datos.get('faltas', 0)
    incompletos = datos.get('registros_incompletos', 0)
    if faltas == 0 and incompletos == 0:
        return "Excelente - Cumplimiento perfecto"
    elif faltas <= 1 and incompletos <= 2:
        return "Bueno - Cumplimiento satisfactorio"
    elif faltas <= 3 and incompletos <= 4:
        return "Regular - Requiere atención"
    else:
        return "Deficiente - Requiere intervención"


# =========================================================
# DIÁLOGO DETALLE COMPLETO (idéntico al original)
# =========================================================
class DialogoDetalleOptimizado(QDialog):
    def __init__(self, parent, datos_informe):
        super().__init__(parent)
        self.datos_informe = datos_informe
        self.setWindowTitle("Detalle Completo de Asistencia")
        self.setGeometry(200, 100, 800, 650)
        self.setStyleSheet("background-color: white; color: #333;")
        self._crear_interfaz()

    def _crear_interfaz(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(15)
        layout.setContentsMargins(20, 20, 20, 20)

        titulo = QLabel(f"Detalle de Asistencia - {self.datos_informe['nombre']}")
        titulo.setStyleSheet("font-size: 24px; font-weight: bold; color: #333;")
        titulo.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(titulo)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet("QScrollArea { border: none; background-color: white; color: #333; }")

        contenedor = QWidget()
        contenedor.setStyleSheet("background-color: white; color: #333;")
        cl = QVBoxLayout(contenedor)
        cl.setSpacing(15)

        self._seccion_info_general(cl)
        self._seccion_estadisticas(cl)
        self._seccion_evaluacion(cl)
        if self.datos_informe.get('detalle_dias'):
            self._seccion_detalle_dias(cl)

        scroll.setWidget(contenedor)
        layout.addWidget(scroll)

        btn_row = QHBoxLayout()
        btn_exp = QPushButton("Exportar datos")
        btn_exp.setStyleSheet("""
            QPushButton { background-color: rgb(15, 110, 86); color: white; font-family: Arial; font-size: 17px; font-weight: bold; padding: 14px 32px; border-radius: 10px; border: none; } QPushButton:hover { background-color: rgb(10, 80, 62); }
        """)
        btn_exp.clicked.connect(self._exportar)
        btn_row.addWidget(btn_exp)
        btn_row.addStretch()
        btn_cerrar = QPushButton("Cerrar")
        btn_cerrar.setStyleSheet("""
            QPushButton { background-color: rgb(163, 45, 45); color: white; font-family: Arial; font-size: 17px; font-weight: bold; padding: 14px 32px; border-radius: 10px; border: none; } QPushButton:hover { background-color: rgb(120,30,30); }
        """)
        btn_cerrar.clicked.connect(self.close)
        btn_row.addWidget(btn_cerrar)
        layout.addLayout(btn_row)

    def _frame_gris(self):
        f = QFrame()
        f.setStyleSheet("""
            QFrame { background-color: #f5f5f5; border: 1px solid #ddd;
                     border-radius: 8px; padding: 15px; }
        """)
        return f

    def _tit(self, texto):
        l = QLabel(texto)
        l.setStyleSheet("color: rgb(26, 26, 26); font-family: Arial; font-size: 20px; font-weight: bold; background: transparent;")
        return l

    def _seccion_info_general(self, layout):
        f  = self._frame_gris()
        fl = QVBoxLayout(f)
        fl.addWidget(self._tit("INFORMACION GENERAL"))
        lbl = QLabel(
            f"ID: {self.datos_informe['id']}\n"
            f"Departamento: {self.datos_informe['departamento']}\n"
            f"Periodo analizado: {self.datos_informe['periodo']}"
        )
        lbl.setStyleSheet("font-size: 20px; color: #333; background: none; border: none;")
        fl.addWidget(lbl)
        layout.addWidget(f)

    def _seccion_estadisticas(self, layout):
        f  = self._frame_gris()
        fl = QVBoxLayout(f)
        fl.addWidget(self._tit("ESTADISTICAS DEL PERIODO"))
        grid = QGridLayout()
        grid.setSpacing(10)
        stats = [
            ("Dias trabajados",        str(self.datos_informe['dias_trabajados'])),
            ("Total horas trabajadas",  self.datos_informe['total_horas_periodo']),
            ("Promedio horas por dia",  f"{self.datos_informe['promedio_horas_dia']} hrs"),
            ("Registros incompletos",   str(self.datos_informe['registros_incompletos'])),
            ("Faltas",                  str(self.datos_informe.get('faltas', 0))),
        ]
        for i, (c, v) in enumerate(stats):
            cl = QLabel(c)
            cl.setStyleSheet("font-size: 16px; color: #666; background: none; border: none;")
            grid.addWidget(cl, i, 0)
            vl = QLabel(str(v))
            vl.setStyleSheet("font-size: 16px; font-weight: bold; color: #333; "
                             "background: none; border: none;")
            vl.setAlignment(Qt.AlignmentFlag.AlignRight)
            grid.addWidget(vl, i, 1)
        fl.addLayout(grid)
        layout.addWidget(f)

    def _seccion_evaluacion(self, layout):
        ev    = _generar_evaluacion(self.datos_informe)
        color = ("#4CAF50" if "Excelente" in ev else
                 "#FF9800" if "Bueno"     in ev else "#f44336")
        f  = QFrame()
        f.setStyleSheet(f"QFrame {{ background-color: {color}; "
                        "border-radius: 8px; padding: 15px; }}")
        fl = QVBoxLayout(f)
        for txt in ["EVALUACION GENERAL", ev]:
            l = QLabel(txt)
            l.setStyleSheet("font-size: 17px; font-weight: bold; color: white; "
                            "background: none; border: none;")
            l.setAlignment(Qt.AlignmentFlag.AlignCenter)
            l.setWordWrap(True)
            fl.addWidget(l)
        layout.addWidget(f)

    def _seccion_detalle_dias(self, layout):
        f  = self._frame_gris()
        fl = QVBoxLayout(f)
        fl.addWidget(self._tit("DETALLE POR DIAS"))

        scroll = QScrollArea()
        scroll.setStyleSheet("QScrollArea { background: white; border: 1px solid #ddd; }")
        scroll.setMaximumHeight(300)
        scroll.setWidgetResizable(True)

        tabla = QWidget()
        tl    = QVBoxLayout(tabla)
        tl.setSpacing(0)
        tl.setContentsMargins(0, 0, 0, 0)

        headers = ["Fecha", "Dia", "Entrada", "Salida", "Horas", "Estado"]
        widths  = [90, 70, 70, 70, 70, 100]

        hdr = QHBoxLayout()
        hdr.setSpacing(0)
        for h, w in zip(headers, widths):
            lbl = QLabel(h)
            lbl.setStyleSheet("""
                QLabel { background-color: #2196F3; color: white; font-size: 14px;
                         font-weight: bold; padding: 10px 6px; border: 1px solid #1976D2; }
            """)
            lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
            lbl.setFixedWidth(w)
            hdr.addWidget(lbl)
        tl.addLayout(hdr)

        for i, dia in enumerate(self.datos_informe['detalle_dias'][:15]):
            fila = QHBoxLayout()
            fila.setSpacing(0)
            bg = "white" if i % 2 == 0 else "#f8f8f8"
            vals = [
                dia.get('fecha', '-'),
                (dia.get('dia_semana', '-') or '-')[:3],
                dia.get('entrada') or '--:--',
                dia.get('salida')  or '--:--',
                dia.get('horas_trabajadas') or '-',
                dia.get('estado', '-'),
            ]
            for v, w in zip(vals, widths):
                cl = QLabel(str(v))
                cl.setStyleSheet(f"QLabel {{ background-color: {bg}; color: #333; "
                                 "font-size: 14px; padding: 8px 4px; "
                                 "border: 1px solid #ddd; }}")
                cl.setAlignment(Qt.AlignmentFlag.AlignCenter)
                cl.setFixedWidth(w)
                fila.addWidget(cl)
            tl.addLayout(fila)

        total = len(self.datos_informe['detalle_dias'])
        if total > 15:
            nota = QLabel(f"Mostrando 15 de {total} dias")
            nota.setStyleSheet("font-size: 13px; color: #666; padding: 6px; "
                               "background: none; border: none;")
            nota.setAlignment(Qt.AlignmentFlag.AlignCenter)
            tl.addWidget(nota)

        scroll.setWidget(tabla)
        fl.addWidget(scroll)
        layout.addWidget(f)

    def _exportar(self):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
            from datetime import datetime as dt

            nombre_arch = (
                f"Informe_{self.datos_informe['nombre']}_"
                f"{dt.now().strftime('%Y%m%d')}.xlsx")
            archivo, _ = QFileDialog.getSaveFileName(
                self, "Guardar informe como Excel", nombre_arch,
                "Archivos Excel (*.xlsx);;Todos los archivos (*.*)")
            if not archivo:
                return

            wb = Workbook()
            ws = wb.active
            ws.title = "Informe Rapido"

            fuente_titulo   = Font(name="Arial", bold=True, size=14, color="FFFFFF")
            fuente_header   = Font(name="Arial", bold=True, size=10, color="FFFFFF")
            fuente_label    = Font(name="Arial", bold=True, size=10)
            fuente_normal   = Font(name="Arial", size=10)
            fill_titulo     = PatternFill("solid", fgColor="1E3A5F")
            fill_header     = PatternFill("solid", fgColor="2E6DA4")
            fill_completo   = PatternFill("solid", fgColor="D4EDDA")
            fill_incompleto = PatternFill("solid", fgColor="FFF3CD")
            fill_falta      = PatternFill("solid", fgColor="F8D7DA")
            fill_info       = PatternFill("solid", fgColor="EBF3FB")
            borde  = Border(left=Side(style="thin"), right=Side(style="thin"),
                            top=Side(style="thin"),  bottom=Side(style="thin"))
            centro = Alignment(horizontal="center", vertical="center")
            izq    = Alignment(horizontal="left",   vertical="center")

            ws.merge_cells("A1:F1")
            ws["A1"] = "INFORME RAPIDO DE ASISTENCIA"
            ws["A1"].font = fuente_titulo
            ws["A1"].fill = fill_titulo
            ws["A1"].alignment = centro
            ws.row_dimensions[1].height = 28

            ev = _generar_evaluacion(self.datos_informe)
            info = [
                ("Empleado",              self.datos_informe['nombre']),
                ("ID",                    self.datos_informe['id']),
                ("Departamento",          self.datos_informe['departamento']),
                ("Periodo",               self.datos_informe['periodo']),
                ("Dias trabajados",       self.datos_informe['dias_trabajados']),
                ("Total horas",           self.datos_informe['total_horas_periodo']),
                ("Promedio horas/dia",    self.datos_informe['promedio_horas_dia']),
                ("Registros incompletos", self.datos_informe['registros_incompletos']),
                ("Faltas",                self.datos_informe.get('faltas', 0)),
                ("Evaluacion",            ev),
            ]
            for fila_i, (label, valor) in enumerate(info, start=2):
                ws.merge_cells(f"A{fila_i}:B{fila_i}")
                ws[f"A{fila_i}"] = label
                ws[f"A{fila_i}"].font = fuente_label
                ws[f"A{fila_i}"].fill = fill_info
                ws[f"A{fila_i}"].alignment = izq
                ws.merge_cells(f"C{fila_i}:F{fila_i}")
                ws[f"C{fila_i}"] = valor
                ws[f"C{fila_i}"].font = fuente_normal
                ws[f"C{fila_i}"].alignment = izq

            if self.datos_informe.get('detalle_dias'):
                fila_header = len(info) + 3
                headers = ["Fecha", "Dia", "Entrada", "Salida", "Horas trabajadas", "Estado"]
                for col_i, h in enumerate(headers, start=1):
                    cell = ws.cell(row=fila_header, column=col_i, value=h)
                    cell.font = fuente_header
                    cell.fill = fill_header
                    cell.alignment = centro
                    cell.border = borde

                fill_map = {"Completo": fill_completo, "Incompleto": fill_incompleto, "Falta": fill_falta}
                for idx, dia in enumerate(self.datos_informe['detalle_dias'], start=fila_header + 1):
                    estado = dia.get('estado', '-')
                    fill   = fill_map.get(estado, PatternFill())
                    valores = [
                        dia.get('fecha', '-'),
                        (dia.get('dia_semana', '-') or '-')[:3],
                        dia.get('entrada') or '-',
                        dia.get('salida')  or '-',
                        dia.get('horas_trabajadas') or '-',
                        estado,
                    ]
                    for col_i, val in enumerate(valores, start=1):
                        cell = ws.cell(row=idx, column=col_i, value=val)
                        cell.font = fuente_normal
                        cell.fill = fill
                        cell.alignment = centro
                        cell.border = borde

            for col_i, ancho in enumerate([13, 8, 9, 9, 17, 13], start=1):
                ws.column_dimensions[get_column_letter(col_i)].width = ancho

            wb.save(archivo)
            QMessageBox.information(self.window(), "Exito", f"Informe exportado:\n{archivo}")
        except Exception as e:
            QMessageBox.critical(self.window(), "Error", f"Error al exportar:\n{str(e)}")


# =========================================================
# WIDGET PRINCIPAL
# =========================================================
class InterfazAsistenciaMejorada(QWidget):
  
    volver_menu = pyqtSignal()

    def __init__(self, sesion=None, parent=None):
        super().__init__(parent)
        self.sesion    = sesion
        self.graphics  = InterfaceGraphics() if GRAPHICS_UTILITY_AVAILABLE else None
        self.procesador = ProcesadorReporteAsistencia() if PROCESSOR_AVAILABLE else None

        self.empleados_disponibles = []
        self.datos_informe: Optional[dict] = None
        self._hilo = None

        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(0, 0, 0, 0)
        main_layout.setSpacing(0)

        self.canvas_fondo = CanvasFondo(self, self.graphics)
        main_layout.addWidget(self.canvas_fondo)

        self.contenedor_widgets = QWidget(self)
        self.contenedor_widgets.setAttribute(Qt.WidgetAttribute.WA_TranslucentBackground)
        self.contenedor_widgets.setStyleSheet("background: transparent;")

        self.timer = QTimer()
        self.timer.timeout.connect(self.canvas_fondo.update)
        self.timer.start(33)

        self.cambiar_pantalla("formulario")

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self.contenedor_widgets.setGeometry(0, 0, self.width(), self.height())

    # ----------------------------------------------------------
    def cambiar_pantalla(self, pantalla: str):
        if self.contenedor_widgets.layout():
            QWidget().setLayout(self.contenedor_widgets.layout())
        for child in self.contenedor_widgets.findChildren(QWidget):
            child.setParent(None)

        layout = QVBoxLayout(self.contenedor_widgets)
        layout.setContentsMargins(60, 50, 60, 50)
        layout.setSpacing(28)

        if pantalla == "formulario":
            self._crear_pantalla_formulario(layout)
            self._cargar_empleados_async()
        elif pantalla == "informe":
            self._crear_pantalla_informe(layout)

    # ----------------------------------------------------------
    # PANTALLA FORMULARIO
    # ----------------------------------------------------------
    def _crear_pantalla_formulario(self, layout):
        layout.addStretch(1)

        titulo = QLabel("Informe rapido de asistencia")
        titulo.setStyleSheet(
            "color: rgb(24, 95, 165); font-family: Arial; font-size: 60px; font-weight: bold; background: transparent;"
        )
        titulo.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(titulo)

        layout.addStretch(1)

        instrucciones = QLabel(
            "Selecciona el empleado del que deseas generar un informe rapido"
        )
        instrucciones.setStyleSheet(
            "color: rgb(26, 26, 26); font-family: Arial; font-size: 25px; background: transparent;"
        )
        instrucciones.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(instrucciones)

        # ComboBox
        self.combo_empleado = QComboBox()
        self.combo_empleado.addItem("-- Cargando empleados... --")
        self.combo_empleado.setStyleSheet(STYLE_COMBO)
        combo_row = QHBoxLayout()
        combo_row.addStretch()
        combo_row.addWidget(self.combo_empleado)
        combo_row.addStretch()
        layout.addLayout(combo_row)

        layout.addSpacing(20)

        # Período
        periodo_lbl = QLabel("Periodo de analisis:")
        periodo_lbl.setStyleSheet(
            "color: rgb(26, 26, 26); font-family: Arial; font-size: 25px; background: transparent;"
        )
        periodo_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        layout.addWidget(periodo_lbl)

        fecha_row = QHBoxLayout()
        fecha_row.addStretch()

        col_ini = QVBoxLayout()
        lbl_ini = QLabel("Fecha inicio")
        lbl_ini.setStyleSheet("color: rgb(26, 26, 26); font-family: Arial; font-size: 20px; font-weight: bold; background: transparent;")
        lbl_ini.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.fecha_inicio_edit = QDateEdit()
        self.fecha_inicio_edit.setCalendarPopup(True)
        self.fecha_inicio_edit.setDate(QDate.currentDate().addDays(-7))
        self.fecha_inicio_edit.setStyleSheet(STYLE_DATE)
        col_ini.addWidget(lbl_ini)
        col_ini.addWidget(self.fecha_inicio_edit)
        fecha_row.addLayout(col_ini)
        fecha_row.addSpacing(30)

        col_fin = QVBoxLayout()
        lbl_fin = QLabel("Fecha fin")
        lbl_fin.setStyleSheet("color: rgb(26, 26, 26); font-family: Arial; font-size: 20px; font-weight: bold; background: transparent;")
        lbl_fin.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.fecha_fin_edit = QDateEdit()
        self.fecha_fin_edit.setCalendarPopup(True)
        self.fecha_fin_edit.setDate(QDate.currentDate())
        self.fecha_fin_edit.setStyleSheet(STYLE_DATE)
        col_fin.addWidget(lbl_fin)
        col_fin.addWidget(self.fecha_fin_edit)
        fecha_row.addLayout(col_fin)
        fecha_row.addStretch()
        layout.addLayout(fecha_row)

        layout.addStretch(2)

        btn_row = QHBoxLayout()
        btn_menu = QPushButton("Menu principal")
        btn_menu.setStyleSheet("QPushButton { background-color: rgb(255, 255, 255); color: rgb(24, 95, 165); font-family: Arial; font-size: 17px; font-weight: bold; padding: 14px 32px; border-radius: 10px; border: 2px solid rgb(24, 95, 165); } QPushButton:hover { background-color: rgb(230, 241, 251); }")
        btn_menu.setCursor(Qt.CursorShape.PointingHandCursor)
        btn_menu.clicked.connect(self._volver)
        btn_row.addWidget(btn_menu)
        btn_row.addStretch()
        self.btn_generar = QPushButton("Generar informe")
        self.btn_generar.setStyleSheet(STYLE_BTN_PRIMARY)
        self.btn_generar.setCursor(Qt.CursorShape.PointingHandCursor)
        self.btn_generar.clicked.connect(self._generar_informe)
        btn_row.addWidget(self.btn_generar)
        layout.addLayout(btn_row)

    def _cargar_empleados_async(self):
        """Lanza carga de empleados desde device_users sin filtro de fechas."""
        if not self.procesador:
            return
        self._hilo = HiloCargaEmpleados(self.procesador)
        self._hilo.completado.connect(self._on_empleados_cargados)
        self._hilo.error.connect(lambda e: self._combo_error())
        self._hilo.start()

    def _on_empleados_cargados(self, empleados: list):
        self.empleados_disponibles = empleados
        if not hasattr(self, 'combo_empleado'):
            return
        self.combo_empleado.clear()
        self.combo_empleado.addItem("-- Seleccionar empleado --")
        for emp in sorted(empleados, key=lambda e: e['nombre']):
            self.combo_empleado.addItem(emp['display_name'])
        if not empleados:
            self.combo_empleado.addItem("Sin empleados en el periodo")

    def _combo_error(self):
        if hasattr(self, 'combo_empleado'):
            self.combo_empleado.clear()
            self.combo_empleado.addItem("Error al cargar empleados")

    # ----------------------------------------------------------
    def _generar_informe(self):
        if not self.procesador:
            QMessageBox.warning(self.window(), "Error", "Modulo de procesamiento no disponible.")
            return

        seleccion = self.combo_empleado.currentText()
        if (not seleccion or seleccion.startswith("--") or
                seleccion.startswith("Sin") or seleccion.startswith("Error")):
            QMessageBox.warning(self.window(), "Advertencia", "Por favor selecciona un empleado.")
            return

        try:
            id_empleado = seleccion.split("ID: ")[1].rstrip(")")
        except IndexError:
            id_empleado = seleccion

        f_inicio = self.fecha_inicio_edit.date().toPyDate()
        f_fin    = self.fecha_fin_edit.date().toPyDate()

        if f_inicio > f_fin:
            QMessageBox.warning(self.window(), "Advertencia",
                                "La fecha de inicio debe ser anterior a la fecha fin.")
            return

        fecha_inicio_dt = datetime.combine(f_inicio, datetime.min.time())
        fecha_fin_dt    = datetime.combine(f_fin,    datetime.max.time())

        self.btn_generar.setEnabled(False)
        self.btn_generar.setText("Consultando...")

        self._hilo = HiloGenerarInforme(
            self.procesador, id_empleado, fecha_inicio_dt, fecha_fin_dt)
        self._hilo.completado.connect(self._on_informe_listo)
        self._hilo.error.connect(self._on_informe_error)
        self._hilo.start()

    def _on_informe_listo(self, datos):
        if hasattr(self, 'btn_generar'):
            self.btn_generar.setEnabled(True)
            self.btn_generar.setText("Generar informe")
        if datos:
            self.datos_informe = datos
            self.cambiar_pantalla("informe")
        else:
            QMessageBox.warning(self.window(), "Sin datos",
                                "No se encontraron registros para este empleado en el periodo.")

    def _on_informe_error(self, msg):
        if hasattr(self, 'btn_generar'):
            self.btn_generar.setEnabled(True)
            self.btn_generar.setText("Generar informe")
        QMessageBox.critical(self.window(), "Error", f"Error al generar el informe:\n{msg}")

    # ----------------------------------------------------------
    # PANTALLA INFORME
    # ----------------------------------------------------------
    def _crear_pantalla_informe(self, layout):
        if not self.datos_informe:
            return

        datos = self.datos_informe

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet("""
            QScrollArea { background: rgb(237, 245, 252); border: none; } QScrollBar:vertical { background: rgb(181, 212, 244); width: 8px; border-radius: 4px; } QScrollBar::handle:vertical { background: rgb(24, 95, 165); border-radius: 4px; }
        """)

        contenedor = QWidget()
        contenedor.setStyleSheet("background: rgb(237, 245, 252);")
        cl = QVBoxLayout(contenedor)
        cl.setSpacing(20)

        # Título
        tit = QLabel("Informe rapido de asistencia")
        tit.setStyleSheet("color: rgb(24, 95, 165); font-family: Arial; font-size: 28px; font-weight: bold; background: transparent;")
        tit.setAlignment(Qt.AlignmentFlag.AlignCenter)
        cl.addWidget(tit)

        # Info empleado
        info = QLabel(
            f"Empleado: {datos['nombre']}\n"
            f"ID: {datos['id']}  |  Departamento: {datos['departamento']}\n"
            f"Periodo: {datos['periodo']}"
        )
        info.setStyleSheet("color: rgb(26, 26, 26); font-family: Arial; font-size: 20px; background: transparent;")
        info.setAlignment(Qt.AlignmentFlag.AlignCenter)
        cl.addWidget(info)

        cl.addSpacing(10)

        resumen_lbl = QLabel("Resumen del periodo:")
        resumen_lbl.setStyleSheet("color: rgb(26, 26, 26); font-family: Arial; font-size: 20px; font-weight: bold; background: transparent;")
        resumen_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        cl.addWidget(resumen_lbl)

        # Tarjeta estadísticas
        stats_frame = QFrame()
        stats_frame.setStyleSheet("""
            QFrame { background-color: rgb(255, 255, 255); border-radius: 12px; border: 1px solid rgb(181, 212, 244); padding: 20px; }
        """)
        sl = QVBoxLayout(stats_frame)
        sl.setSpacing(8)

        for concepto, valor in [
            ("Horas trabajadas:",      datos['total_horas_periodo']),
            ("Dias trabajados:",       datos['dias_trabajados']),
            ("Registros incompletos:", datos['registros_incompletos']),
            ("Faltas:",                datos.get('faltas', 0)),
            ("Promedio horas/dia:",    f"{datos['promedio_horas_dia']} hrs"),
        ]:
            fila = QHBoxLayout()
            cl_lbl = QLabel(concepto)
            cl_lbl.setStyleSheet("color: rgb(26, 26, 26); font-family: Arial; font-size: 20px; background: transparent;")
            fila.addWidget(cl_lbl)
            fila.addStretch()
            vl = QLabel(str(valor))
            vl.setStyleSheet("color: rgb(26, 26, 26); font-family: Arial; font-size: 18px; font-weight: bold; background: transparent;")
            fila.addWidget(vl)
            sl.addLayout(fila)

        cl.addWidget(stats_frame)

        # Evaluación
        ev = _generar_evaluacion(datos)
        ev_lbl = QLabel(f"Evaluacion: {ev}")
        ev_lbl.setStyleSheet("color: rgb(26, 26, 26); font-family: Arial; font-size: 20px; font-weight: bold; background: transparent;")
        ev_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        ev_lbl.setWordWrap(True)
        cl.addWidget(ev_lbl)

        cl.addStretch()
        scroll.setWidget(contenedor)
        layout.addWidget(scroll)

        # Botones
        btn_row = QHBoxLayout()

        btn_nueva = QPushButton("Nueva consulta")
        btn_nueva.setStyleSheet("QPushButton { background-color: rgb(255, 255, 255); color: rgb(24, 95, 165); font-family: Arial; font-size: 17px; font-weight: bold; padding: 14px 32px; border-radius: 10px; border: 2px solid rgb(24, 95, 165); } QPushButton:hover { background-color: rgb(230, 241, 251); }")
        btn_nueva.setCursor(Qt.CursorShape.PointingHandCursor)
        btn_nueva.clicked.connect(lambda: self.cambiar_pantalla("formulario"))
        btn_row.addWidget(btn_nueva)

        btn_row.addStretch()

        btn_detalle = QPushButton("Ver detalle completo")
        btn_detalle.setStyleSheet("QPushButton { background-color: rgb(15, 110, 86); color: white; font-family: Arial; font-size: 17px; font-weight: bold; padding: 14px 32px; border-radius: 10px; border: none; } QPushButton:hover { background-color: rgb(10, 80, 62); }")
        btn_detalle.setCursor(Qt.CursorShape.PointingHandCursor)
        btn_detalle.clicked.connect(self._mostrar_detalle)
        btn_row.addWidget(btn_detalle)

        btn_row.addStretch()

        btn_menu = QPushButton("Menu principal")
        btn_menu.setStyleSheet("QPushButton { background-color: rgb(255, 255, 255); color: rgb(24, 95, 165); font-family: Arial; font-size: 17px; font-weight: bold; padding: 14px 32px; border-radius: 10px; border: 2px solid rgb(24, 95, 165); } QPushButton:hover { background-color: rgb(230, 241, 251); }")
        btn_menu.setCursor(Qt.CursorShape.PointingHandCursor)
        btn_menu.clicked.connect(self._volver)
        btn_row.addWidget(btn_menu)

        layout.addLayout(btn_row)

    def _mostrar_detalle(self):
        if self.datos_informe:
            DialogoDetalleOptimizado(self, self.datos_informe).exec()

    def _volver(self):
        self.timer.stop()
        self.volver_menu.emit()


# =========================================================
# MAIN (debug standalone)
# =========================================================
def main():
    if os.name == 'nt' and GRAPHICS_UTILITY_AVAILABLE:
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(2)
        except Exception:
            pass

    ocultar_consola_windows()
    app = QApplication(sys.argv)

    class SesionDebug:
        user_name = "admin"

    ventana = QMainWindow()
    ventana.setWindowTitle("Consulta Rapida - Debug")
    widget = InterfazAsistenciaMejorada(sesion=SesionDebug())
    widget.volver_menu.connect(app.quit)
    ventana.setCentralWidget(widget)
    ventana.showMaximized()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()